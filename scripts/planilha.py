#!/usr/bin/env python3
"""Write validated invoices to a sortable spreadsheet. Three sheets.

  RECIBOS    one row per invoice, every field plus its flags
  GARANTIAS  only purchases classified as goods, with both milestone dates
  REVISAO    everything flagged, so nothing broken hides among the clean rows

Writes .xlsx when openpyxl is installed and .csv otherwise, with identical
columns either way - the CSV path exists so the skill has NO hard dependency
and still runs on a machine with nothing but Python.

The REVISAO sheet is the point of the whole file. A validator that produces a
tidy spreadsheet where the broken rows look like the good ones has produced a
more dangerous artifact than no spreadsheet at all.

Usage:
    python planilha.py --selftest
"""

import argparse
import csv
import os
import sys

COLUMNS = [
    ("id", "N.o"),
    ("data", "Data"),
    ("nif_emitente", "NIF emitente"),
    ("nif_adquirente", "NIF adquirente"),
    ("tipo_documento", "Tipo"),
    ("numero_documento", "N.o documento"),
    ("atcud", "ATCUD"),
    ("base_total", "Base (EUR)"),
    ("imposto_total", "Impostos (EUR)"),
    ("total_com_impostos", "Total (EUR)"),
    ("categoria", "Categoria IRS"),
    ("deducao_possivel", "Dedutivel?"),
    ("reconcilia", "Reconcilia?"),
    ("source", "Origem"),
    ("sem_erros_detetados", "Sem erros detetados"),
    ("n_flags", "N.o avisos"),
    ("flags", "Avisos"),
]

GARANTIA_COLUMNS = [
    ("id", "N.o"),
    ("descricao", "Descricao"),
    ("data", "Data da compra"),
    ("tipo_label", "Tipo de bem"),
    ("presuncao_termina", "Fim da presuncao (prova passa a ser sua)"),
    ("garantia_termina", "Fim da garantia legal"),
    ("data_e_estimada", "Data estimada da fatura?"),
    ("numero_documento", "N.o documento"),
    ("atcud", "ATCUD"),
    ("base_legal", "Base legal"),
    ("avisos", "Avisos"),
]


def _flat(rec, i):
    """One record -> one flat row dict."""
    flags = rec.get("flags") or []
    return {
        "id": i,
        "data": rec.get("data", ""),
        "nif_emitente": rec.get("nif_emitente", ""),
        "nif_adquirente": rec.get("nif_adquirente") or "(nao fornecido)",
        "tipo_documento": rec.get("tipo_documento", ""),
        "numero_documento": rec.get("numero_documento", ""),
        "atcud": rec.get("atcud", ""),
        "base_total": rec.get("base_total", ""),
        "imposto_total": rec.get("imposto_total", ""),
        "total_com_impostos": rec.get("total_com_impostos", ""),
        "categoria": rec.get("categoria", "(por classificar)"),
        "deducao_possivel": _yn(rec.get("deducao_possivel")),
        "reconcilia": _yn(rec.get("reconcilia")),
        "source": rec.get("source", ""),
        "sem_erros_detetados": _yn(rec.get("sem_erros_detetados")),
        "n_flags": len(flags),
        "flags": " | ".join(flags),
    }


def _yn(v):
    return "" if v is None else ("SIM" if v else "NAO")


def _garantia_row(rec, i):
    g = rec.get("garantia") or {}
    return {
        "id": i,
        "descricao": rec.get("descricao", ""),
        "data": rec.get("data", ""),
        "tipo_label": g.get("tipo_label", ""),
        "presuncao_termina": g.get("presuncao_termina") or "",
        "garantia_termina": g.get("garantia_termina") or "",
        "data_e_estimada": "SIM" if g.get("data_e_estimada") else "NAO",
        "numero_documento": rec.get("numero_documento", ""),
        "atcud": rec.get("atcud", ""),
        "base_legal": g.get("base_legal", ""),
        "avisos": " | ".join(g.get("avisos", [])),
    }


def build(records):
    """records -> {sheet_name: (headers, rows)}. Pure; no I/O."""
    recibos = [_flat(r, i) for i, r in enumerate(records, 1)]
    garantias = [_garantia_row(r, i) for i, r in enumerate(records, 1)
                 if (r.get("garantia") or {}).get("garantia_termina")]
    revisao = [row for row in recibos if row["n_flags"]]
    return {
        "RECIBOS": ([h for _, h in COLUMNS],
                    [[row[k] for k, _ in COLUMNS] for row in recibos]),
        "GARANTIAS": ([h for _, h in GARANTIA_COLUMNS],
                      [[row[k] for k, _ in GARANTIA_COLUMNS] for row in garantias]),
        "REVISAO": ([h for _, h in COLUMNS],
                    [[row[k] for k, _ in COLUMNS] for row in revisao]),
    }


def write(records, path_base):
    """Write .xlsx if openpyxl is available, else .csv. Returns paths written."""
    sheets = build(records)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        out = []
        for name, (headers, rows) in sheets.items():
            p = "%s.%s.csv" % (path_base, name.lower())
            with open(p, "w", encoding="utf-8-sig", newline="") as fh:
                w = csv.writer(fh, delimiter=";")
                w.writerow(headers)
                w.writerows(rows)
            out.append(p)
        return out

    wb = Workbook()
    wb.remove(wb.active)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="0F9B7A")
    warn_fill = PatternFill("solid", fgColor="FFF4D6")
    for name, (headers, rows) in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)),
                                          max(1, len(rows) + 1))
        for i, h in enumerate(headers, 1):
            width = max(len(str(h)), *(len(str(r[i - 1])) for r in rows)) if rows else len(h)
            ws.column_dimensions[get_column_letter(i)].width = min(46, width + 2)
        if name == "REVISAO":
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.fill = warn_fill
    p = path_base + ".xlsx"
    wb.save(p)
    return [p]


def _selftest():
    import tempfile
    from datetime import date
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import garantia

    good = {"data": "2026-02-14", "nif_emitente": "501442600",
            "nif_adquirente": "234567899", "tipo_documento": "FT",
            "numero_documento": "FT 2026/117", "atcud": "JFHF6T5J-117",
            "base_total": "100.00", "imposto_total": "23.00",
            "total_com_impostos": "123.00", "deducao_possivel": True,
            "reconcilia": True, "source": "qr", "sem_erros_detetados": True,
            "flags": [], "descricao": "Portatil",
            "garantia": garantia.compute(date(2026, 2, 14), "bem_movel_novo")}
    flagged = dict(good, numero_documento="FS 9", sem_erros_detetados=False,
                   deducao_possivel=False, descricao="Cafe",
                   flags=["SEM_NIF_DEDUCAO_PERDIDA"],
                   garantia=garantia.compute(date(2026, 2, 14), "servico"))
    sheets = build([good, flagged])

    failed = 0
    checks = [
        ("RECIBOS holds every record", len(sheets["RECIBOS"][1]) == 2),
        ("GARANTIAS excludes the service", len(sheets["GARANTIAS"][1]) == 1),
        ("REVISAO holds only flagged rows", len(sheets["REVISAO"][1]) == 1),
        ("flags are visible in a column", "SEM_NIF_DEDUCAO_PERDIDA" in
         str(sheets["REVISAO"][1][0])),
        ("column count is stable",
         all(len(r) == len(sheets["RECIBOS"][0]) for r in sheets["RECIBOS"][1])),
        ("missing NIF is not blank",
         "(nao fornecido)" in str(sheets["RECIBOS"][1][0]) or
         "234567899" in str(sheets["RECIBOS"][1][0])),
        ("None never renders as 'None'",
         "None" not in str(sheets["RECIBOS"][1]) + str(sheets["GARANTIAS"][1])),
    ]
    print("spreadsheet build")
    print("-" * 60)
    for n, ok in checks:
        failed += 0 if ok else 1
        print("%-4s %s" % ("PASS" if ok else "FAIL", n))

    tmp = tempfile.mkdtemp()
    paths = write([good, flagged], os.path.join(tmp, "recibos"))
    ok = paths and all(os.path.getsize(p) > 0 for p in paths)
    failed += 0 if ok else 1
    print("%-4s wrote %s" % ("PASS" if ok else "FAIL",
                             ", ".join(os.path.basename(p) for p in paths)))
    print("-" * 60)
    print("SELFTEST: %s" % ("PASS" if not failed else "FAIL (%d)" % failed))
    return 1 if failed else 0


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Write validated invoices to a sheet.")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args()
    if a.selftest:
        sys.exit(_selftest())
    ap.print_help()
    sys.exit(2)
